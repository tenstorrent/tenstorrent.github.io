#!/usr/bin/env python3
"""Assemble TT-QuietBox 2 User Guide DOCX from the Product Doc template.

Fills the Cursor QB2 template with Specifications → Setup → Compliance,
preserving template title styles, legal/trademark front matter, and footers.
"""

from __future__ import annotations

import os
import re
import sys
from copy import deepcopy
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parents[1]
SRC = REPO / "core" / "systems" / "quietbox" / "quietbox-bh-2"
# Approved style shell — regenerate content into this look going forward.
GOLDEN = SRC / "for-cursor-tt-quietbox-2-user-guide.docx"
TEMPLATE = SRC / "_product-doc-template.docx"
OUT_DOCX = SRC / "tt-quietbox-2-user-guide.docx"
OUT_PDF = SRC / "tt-quietbox-2-user-guide.pdf"

CHAPTERS = [
    ("specifications.md", "Specifications"),
    ("setup.md", "Hardware and Software Setup"),
    ("compliance-qb2.md", "Compliance and Legal"),
]

SOURCE_MD = [SRC / name for name, _ in CHAPTERS]
BUILD_INPUTS = [
    *SOURCE_MD,
    GOLDEN,
    Path(__file__).resolve(),
]

IMG_WIDTH = Inches(4.5)

# Numbering instances preserved in the golden template (word/numbering.xml).
BULLET_NUM_ID = "8"   # •
# Decimal abstract used when allocating a fresh numId per ordered list (restarts at 1).
DECIMAL_ABSTRACT_NUM_ID = "22"  # numId 14 -> abstract 22

# Callout fills from the golden DOCX.
CALLOUT_COLORS = {
    "caution": {"title": "FFC000", "body": "FFEECD"},
    "danger": {"title": "FF9E8A", "body": "F5D8D2"},
    "important": {"title": "F79646", "body": "FFEECD"},
    "note": {"title": "548DD4", "body": "DBE5F1"},
}
TABLE_HEADER_FILL = "EEECE1"
PARA_SPACE_AFTER = Pt(10)


def delete_paragraph(paragraph: Paragraph) -> None:
    el = paragraph._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def delete_table(table) -> None:
    el = table._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def set_runs_text(paragraph: Paragraph, text: str) -> None:
    """Replace paragraph text while keeping the first run's formatting if possible."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_paragraph_shading(paragraph: Paragraph, fill: str) -> None:
    """Apply a solid paragraph background fill (e.g. 'FFC000')."""
    pPr = paragraph._p.get_or_add_pPr()
    shd = pPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        pPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_paragraph_numbering(paragraph: Paragraph, num_id: str, ilvl: str = "0") -> None:
    """Attach a numbering definition so List Paragraph shows bullets/numbers."""
    pPr = paragraph._p.get_or_add_pPr()
    numPr = pPr.find(qn("w:numPr"))
    if numPr is not None:
        pPr.remove(numPr)
    numPr = OxmlElement("w:numPr")
    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), ilvl)
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), num_id)
    numPr.append(ilvl_el)
    numPr.append(num_id_el)
    pPr.append(numPr)


def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_table_borders(table) -> None:
    """Force visible borders on all sides (matches golden template)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def apply_space_after(paragraph: Paragraph, pts: Pt = PARA_SPACE_AFTER) -> None:
    paragraph.paragraph_format.space_after = pts


def set_paragraph_indent(
    paragraph: Paragraph, left_twips: int, hanging_twips: int = 360
) -> None:
    """Override list/body left indent (twips)."""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:left"), str(left_twips))
    ind.set(qn("w:hanging"), str(hanging_twips))


def set_exact_line_spacing(paragraph: Paragraph, pts: float) -> None:
    """Set exact line spacing in points (for thin colored spacers)."""
    pPr = paragraph._p.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), str(int(pts * 20)))  # twips
    spacing.set(qn("w:lineRule"), "exact")


def add_shaded_spacer(doc: Document, fill: str, style: str, pts: float = 3.0) -> None:
    """Colored empty padding strip inside a callout box."""
    p = doc.add_paragraph(style=style)
    set_paragraph_shading(p, fill)
    set_exact_line_spacing(p, pts)
    if p.runs:
        p.runs[0].text = ""
    else:
        p.add_run("")


def _ensure_vert_align_superscript(rPr) -> None:
    vert = rPr.find(qn("w:vertAlign"))
    if vert is None:
        vert = OxmlElement("w:vertAlign")
        rPr.append(vert)
    vert.set(qn("w:val"), "superscript")


def append_text_with_trademarks(paragraph: Paragraph, text: str, *, bold: bool | None = None) -> None:
    """Append text, putting each ® in its own superscript run."""
    if not text:
        return
    parts = text.split("®")
    for i, part in enumerate(parts):
        if part:
            run = paragraph.add_run(part)
            if bold is not None:
                run.bold = bold
        if i < len(parts) - 1:
            reg = paragraph.add_run("®")
            reg.font.superscript = True
            if bold is not None:
                reg.bold = bold


def superscript_registered_in_paragraph(paragraph: Paragraph) -> None:
    """Ensure every ® in a paragraph is a superscript run."""
    for run in list(paragraph.runs):
        text = run.text
        if not text or "®" not in text:
            continue
        if text == "®":
            run.font.superscript = True
            rPr = run._element.get_or_add_rPr()
            _ensure_vert_align_superscript(rPr)
            continue

        parts = text.split("®")
        run.text = parts[0]
        anchor = run._element
        base_rPr = anchor.find(qn("w:rPr"))

        for part in parts[1:]:
            r_reg = OxmlElement("w:r")
            rPr = deepcopy(base_rPr) if base_rPr is not None else OxmlElement("w:rPr")
            _ensure_vert_align_superscript(rPr)
            r_reg.append(rPr)
            t_el = OxmlElement("w:t")
            t_el.text = "®"
            r_reg.append(t_el)
            anchor.addnext(r_reg)
            anchor = r_reg

            if part:
                r_rest = OxmlElement("w:r")
                if base_rPr is not None:
                    r_rest.append(deepcopy(base_rPr))
                t_rest = OxmlElement("w:t")
                if part.startswith(" ") or part.endswith(" "):
                    t_rest.set(qn("xml:space"), "preserve")
                t_rest.text = part
                r_rest.append(t_rest)
                anchor.addnext(r_rest)
                anchor = r_rest


def superscript_all_registered(doc: Document) -> None:
    """Superscript ® in body, tables, headers, and footers."""
    for paragraph in doc.paragraphs:
        superscript_registered_in_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    superscript_registered_in_paragraph(paragraph)
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            try:
                for paragraph in part.paragraphs:
                    superscript_registered_in_paragraph(paragraph)
            except Exception:  # noqa: BLE001
                continue


def configure_toc_from_specifications(doc: Document) -> None:
    """Point the TOC field at a bookmark starting at Specifications (skip front matter)."""
    for el in doc.element.body.iter(qn("w:instrText")):
        if el.text and "TOC" in el.text:
            el.text = " TOC \\h \\u \\z \\b AppContents "

    body = doc.element.body
    # Remove any prior AppContents bookmarks.
    for start in list(body.iter(qn("w:bookmarkStart"))):
        if start.get(qn("w:name")) != "AppContents":
            continue
        bid = start.get(qn("w:id"))
        parent = start.getparent()
        if parent is not None:
            parent.remove(start)
        for end in list(body.iter(qn("w:bookmarkEnd"))):
            if end.get(qn("w:id")) == bid:
                ep = end.getparent()
                if ep is not None:
                    ep.remove(end)

    spec = None
    for paragraph in doc.paragraphs:
        if (
            paragraph.text.strip() == "Specifications"
            and paragraph.style
            and paragraph.style.name.startswith("Heading")
        ):
            spec = paragraph
            break
    if spec is None:
        return

    used = {
        el.get(qn("w:id"))
        for el in body.iter(qn("w:bookmarkStart"))
        if el.get(qn("w:id")) is not None
    }
    bid = "100"
    while bid in used:
        bid = str(int(bid) + 1)

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), "AppContents")
    spec._p.insert(0, start)

    last_p = None
    for child in body.iterchildren():
        if child.tag in (qn("w:p"), qn("w:tbl"), qn("w:sdt")):
            last_p = child
    if last_p is None:
        last_p = spec._p
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    last_p.addnext(end)


def replace_run_with_page_field(run) -> None:
    """Replace a run's text with a Word PAGE field."""
    run.text = ""
    r = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    # Placeholder shown until Word updates fields
    stub = OxmlElement("w:t")
    stub.text = "1"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r.append(fld_begin)
    r2 = OxmlElement("w:r")
    r2.append(instr)
    r.addnext(r2)
    r3 = OxmlElement("w:r")
    r3.append(fld_sep)
    r2.addnext(r3)
    r4 = OxmlElement("w:r")
    r4.append(stub)
    r3.addnext(r4)
    r5 = OxmlElement("w:r")
    r5.append(fld_end)
    r4.addnext(r5)


def ensure_footer_page_numbers(doc: Document) -> None:
    """Single-line footer: left copyright/title, right-aligned PAGE field."""
    for section in doc.sections:
        footer = section.footer
        if footer.is_linked_to_previous:
            continue

        # Usable width between margins — Word tab positions are in twips
        # (1 inch = 914400 EMUs = 1440 twips ⇒ divide EMUs by 635).
        usable_emu = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
        usable = max(usable_emu // 635, 1)

        # Collect existing left-side text (ignore bare page digits / spacer runs).
        left_bits = []
        for para in footer.paragraphs:
            for run in para.runs:
                text = run.text.replace("\xa0", " ").strip()
                if not text or text.isdigit():
                    continue
                left_bits.append(text)
        left_text = "   ".join(left_bits) if left_bits else (
            "© 2026 Tenstorrent USA, Inc.  https://tenstorrent.com/   "
            "TT-QuietBox 2 (Blackhole) User Guide"
        )

        paras = list(footer.paragraphs)
        para = paras[0] if paras else footer.add_paragraph()
        # Clear existing runs / content on the first paragraph.
        for run in list(para.runs):
            run._element.getparent().remove(run._element)
        for child in list(para._p):
            if child.tag != qn("w:pPr"):
                para._p.remove(child)

        pPr = para._p.get_or_add_pPr()
        old_tabs = pPr.find(qn("w:tabs"))
        if old_tabs is not None:
            pPr.remove(old_tabs)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), str(usable))
        tabs.append(tab)
        pPr.append(tabs)

        spacing = pPr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            pPr.append(spacing)
        spacing.set(qn("w:after"), "0")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")

        left_run = para.add_run(left_text)
        left_run.font.size = Pt(10)
        para.add_run("\t")
        page_run = para.add_run()
        page_run.font.size = Pt(10)
        replace_run_with_page_field(page_run)

        for extra in paras[1:]:
            delete_paragraph(extra)


class NumberingAllocator:
    """Allocate a fresh numId per ordered list so each list restarts at 1."""

    def __init__(self, doc: Document):
        self.doc = doc
        self._next_id = None
        self._numbering = None
        try:
            self._numbering = doc.part.numbering_part._element
            existing = [
                int(n.get(qn("w:numId")))
                for n in self._numbering.findall(qn("w:num"))
                if n.get(qn("w:numId"))
            ]
            self._next_id = (max(existing) + 1) if existing else 100
        except Exception:
            self._next_id = None

    def new_ordered_num_id(self) -> str:
        if self._numbering is None or self._next_id is None:
            return "14"  # fallback: may continue numbering
        num_id = str(self._next_id)
        self._next_id += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), num_id)
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), DECIMAL_ABSTRACT_NUM_ID)
        num.append(abstract)
        # Explicit restart at 1
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        self._numbering.append(num)
        return num_id


_NUMBERING: NumberingAllocator | None = None


def get_numbering() -> NumberingAllocator | None:
    return _NUMBERING


def strip_frontmatter(text: str) -> str:
    if text.startswith("---"):
        end = text.find("\n---", 3)
        if end != -1:
            return text[end + 4 :].lstrip("\n")
    return text


def clean_md(text: str) -> str:
    text = strip_frontmatter(text)
    text = re.sub(r"```\{raw\} html\n.*?```\n?", "", text, flags=re.DOTALL)
    text = re.sub(r"^:class:\s*\w+\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"\(safety-warnings\)=\n?", "", text)
    text = re.sub(r"\{ref\}`([^<]+)<[^>]+>`", r"\1", text)
    text = re.sub(r"<sup>®</sup>", "®", text)
    # Keep <br> markers so markdown tables stay on one line; expand later in cells.
    text = re.sub(r"<br\s*/?>", " <br> ", text, flags=re.IGNORECASE)
    text = re.sub(r"<small><em>(.*?)</em></small>", r"\1", text, flags=re.DOTALL)
    text = re.sub(r"</?small>", "", text)
    text = re.sub(r"</?em>", "", text)
    text = re.sub(r"\\-\s+", "- ", text)
    # Drop site-only "What to do next" HTML hero; keep a short plain note later if present
    text = re.sub(
        r"## What to do next\?\s*\n```\{raw\} html\n.*?```",
        "## What to do next?\n\n"
        "Your TT-QuietBox 2 is ready. Continue with the welcome guide at "
        "https://docs.tenstorrent.com/ for hands-on lessons.\n",
        text,
        flags=re.DOTALL,
    )
    return text


def parse_blocks(md: str) -> list[tuple]:
    """Parse cleaned markdown into simple blocks for python-docx."""
    blocks: list[tuple] = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # MyST / markdown figure
        fig = re.match(r"```\{figure\}\s+(\./\S+)", stripped)
        if fig:
            path = fig.group(1)
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(("image", path))
            continue

        md_img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if md_img:
            blocks.append(("image", md_img.group(2)))
            i += 1
            continue

        if stripped.startswith("```"):
            lang = stripped.strip("`").strip()
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(("code", "\n".join(code_lines), lang))
            continue

        # Admonitions
        adm = re.match(r":::\{([^}]+)\}(.*)$", stripped)
        if adm:
            kind = adm.group(1).strip()
            title = adm.group(2).strip()
            i += 1
            body = []
            while i < len(lines) and lines[i].strip() != ":::":
                if not lines[i].strip().startswith(":class:"):
                    body.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            label = title or kind.split()[0].title()
            blocks.append(("admonition", label, "\n".join(body).strip()))
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                if re.match(r"^\|[\s:-]+\|", tl):
                    continue
                cells = [
                    c.replace("<br>", "\n").replace("<br />", "\n").replace("<br/>", "\n").strip()
                    for c in tl.strip("|").split("|")
                ]
                rows.append(cells)
            if rows:
                blocks.append(("table", rows))
            continue

        heading = re.match(r"^(#{1,6})\s+\*{0,2}(.+?)\*{0,2}\s*$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip().strip("*")
            blocks.append(("heading", level, text))
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            items = []
            ordered = bool(re.match(r"^\d+\.\s+", stripped))
            while i < len(lines):
                s = lines[i].strip()
                if re.match(r"^[-*]\s+", s):
                    items.append(re.sub(r"^[-*]\s+", "", s))
                    ordered = False
                    i += 1
                elif re.match(r"^\d+\.\s+", s):
                    items.append(re.sub(r"^\d+\.\s+", "", s))
                    ordered = True
                    i += 1
                elif s.startswith("```{figure}"):
                    break
                elif s.startswith("```"):
                    break
                elif s.startswith("#"):
                    break
                elif s.startswith("|"):
                    break
                elif s.startswith(":::"):
                    break
                elif not s:
                    # blank line ends list unless next continues list
                    if i + 1 < len(lines) and (
                        re.match(r"^[-*]\s+", lines[i + 1].strip())
                        or re.match(r"^\d+\.\s+", lines[i + 1].strip())
                    ):
                        i += 1
                        continue
                    break
                else:
                    # continuation of previous item
                    if items:
                        items[-1] += " " + s
                        i += 1
                    else:
                        break
            blocks.append(("list", ordered, items))
            continue

        if stripped in ("---", "***"):
            i += 1
            continue

        # Paragraph (may span lines until blank)
        para_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                break
            if nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("```") or nxt.startswith(":::"):
                break
            if re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+\.\s+", nxt):
                break
            if nxt in ("---", "***"):
                break
            para_lines.append(nxt)
            i += 1
        text = " ".join(para_lines)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
        blocks.append(("para", text))

    return blocks


def add_heading(doc: Document, text: str, level: int) -> None:
    level = min(max(level, 1), 4)
    p = doc.add_heading(level=level)
    append_text_with_trademarks(p, text)
    apply_space_after(p)


def add_para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    append_text_with_trademarks(p, text)
    apply_space_after(p)


def add_list(
    doc: Document,
    ordered: bool,
    items: list[str],
    shade: str | None = None,
    num_id: str | None = None,
    left_twips: int | None = None,
) -> str | None:
    """Add a list. For ordered lists, returns the numId used (for continuation)."""
    style_names = {s.name for s in doc.styles}
    style = "List Paragraph" if "List Paragraph" in style_names else "Normal"
    if ordered:
        if num_id is None:
            allocator = get_numbering()
            num_id = allocator.new_ordered_num_id() if allocator else "14"
    else:
        num_id = BULLET_NUM_ID
    for item in items:
        item = re.sub(r"\*\*(.+?)\*\*", r"\1", item)
        item = re.sub(r"`([^`]+)`", r"\1", item)
        item = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", item)
        p = doc.add_paragraph(style=style)
        append_text_with_trademarks(p, item)
        set_paragraph_numbering(p, num_id)
        if left_twips is not None:
            set_paragraph_indent(p, left_twips)
        if shade:
            set_paragraph_shading(p, shade)
        apply_space_after(p)
    return num_id if ordered else None


def add_code(doc: Document, code: str) -> None:
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph(line)
        p.style = doc.styles["Normal"]
        apply_space_after(p)
        for run in p.runs:
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            run.font.size = Pt(9)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    for candidate in ("TableNormal", "Table Grid", "Normal Table"):
        try:
            table.style = candidate
            break
        except Exception:
            continue
    set_table_borders(table)
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell_text = cell_text.replace("<br>", "\n")
            cell_text = re.sub(r"<[^>]+>", "", cell_text)
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = cell_text
            if r_idx == 0:
                set_cell_shading(cell, TABLE_HEADER_FILL)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def add_image(doc: Document, rel_path: str) -> None:
    from io import BytesIO

    from PIL import Image as PILImage

    path = (SRC / rel_path).resolve()
    if not path.exists():
        path = SRC / rel_path.lstrip("./")
    if not path.exists():
        add_para(doc, f"[Missing image: {rel_path}]")
        return
    try:
        im = PILImage.open(path)
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGB")
        max_px = 2000
        if max(im.size) > max_px:
            im.thumbnail((max_px, max_px), PILImage.Resampling.LANCZOS)
        buf = BytesIO()
        if path.suffix.lower() == ".png" or im.mode == "RGBA":
            im.save(buf, format="PNG")
        else:
            im = im.convert("RGB")
            im.save(buf, format="JPEG", quality=85)
        buf.seek(0)
        doc.add_picture(buf, width=IMG_WIDTH)
    except Exception as exc:  # noqa: BLE001
        add_para(doc, f"[Could not insert image {rel_path}: {exc}]")


def _callout_kind(label: str) -> str:
    lower = label.lower()
    if "danger" in lower:
        return "danger"
    if "hot surface" in lower or "caution" in lower:
        return "caution"
    if "important" in lower or "warning" in lower:
        return "important"
    if "note" in lower:
        return "note"
    return "note"


def add_admonition(doc: Document, label: str, body: str) -> None:
    """Match golden callouts: colored title + fully shaded body (including lists)."""
    style_names = {s.name for s in doc.styles}
    no_spacing = "No Spacing" if "No Spacing" in style_names else "Normal"
    kind = _callout_kind(label)
    colors = CALLOUT_COLORS[kind]

    label_clean = re.sub(r"^(admonition\s+)?", "", label, flags=re.I).strip() or "Note"
    lower = label_clean.lower()
    if "hot surface" in lower or lower == "caution":
        title = "⚠ Caution: Hot Surface"
        title_style = no_spacing
    elif lower == "danger":
        title = "⚠ Danger"
        title_style = "Normal"
    elif lower in ("important", "warning"):
        title = "Important"
        title_style = no_spacing
    elif lower == "note":
        title = "Note:"
        title_style = "Normal"
    else:
        title = label_clean if label_clean.endswith(":") else f"{label_clean}:"
        title_style = "Normal"

    # 3pt colored padding above the callout.
    add_shaded_spacer(doc, colors["title"], title_style, pts=3.0)

    title_p = doc.add_paragraph(style=title_style)
    append_text_with_trademarks(title_p, title, bold=True)
    set_paragraph_shading(title_p, colors["title"])
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(0)

    body = body.strip()
    body_style = no_spacing
    if not body:
        add_shaded_spacer(doc, colors["title"], title_style, pts=3.0)
        return
    lines = [ln.strip() for ln in body.splitlines() if ln.strip()]
    list_items = []
    prose = []
    for s in lines:
        if re.match(r"^[-*]\s+", s) or re.match(r"^\d+\.\s+", s):
            list_items.append(re.sub(r"^([-*]|\d+\.)\s+", "", s))
        else:
            prose.append(s)

    for chunk in prose:
        chunk = re.sub(r"\*\*(.+?)\*\*", r"\1", chunk)
        chunk = re.sub(r"`([^`]+)`", r"\1", chunk)
        chunk = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", chunk)
        p = doc.add_paragraph(style=body_style)
        append_text_with_trademarks(p, chunk)
        set_paragraph_shading(p, colors["body"])
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    # Keep callout list items shaded + bulleted like the golden Danger/ESD boxes.
    style_names = {s.name for s in doc.styles}
    list_style = "List Paragraph" if "List Paragraph" in style_names else body_style
    for item in list_items:
        item = re.sub(r"\*\*(.+?)\*\*", r"\1", item)
        item = re.sub(r"`([^`]+)`", r"\1", item)
        item = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", item)
        p = doc.add_paragraph(style=list_style)
        append_text_with_trademarks(p, item)
        set_paragraph_numbering(p, BULLET_NUM_ID)
        set_paragraph_shading(p, colors["body"])
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

    # 3pt colored padding below the callout.
    add_shaded_spacer(doc, colors["body"], body_style, pts=3.0)


def render_blocks(doc: Document, blocks: list[tuple], skip_first_h1: bool = True) -> None:
    skipped_h1 = False
    # Continue a single ordered list across intervening code/image blocks (Step 9 pattern).
    active_ordered_num_id: str | None = None
    current_h2: str | None = None
    for block in blocks:
        kind = block[0]
        if kind == "heading":
            _, level, text = block
            if skip_first_h1 and not skipped_h1 and level == 1:
                skipped_h1 = True
                continue
            active_ordered_num_id = None
            if level == 1:
                current_h2 = None
            elif level == 2:
                current_h2 = text
            add_heading(doc, text, level)
        elif kind == "para":
            active_ordered_num_id = None
            add_para(doc, block[1])
        elif kind == "list":
            ordered = block[1]
            items = block[2]
            # Hardware/setup ordered lists sit two half-inch indents too far right by default.
            left_twips = None
            if current_h2:
                step = current_h2.lower()
                if any(
                    step.startswith(prefix)
                    for prefix in (
                        "step 2",
                        "step 4",
                        "step 7",
                        "step 8",
                        "step 9",
                    )
                ):
                    left_twips = 720  # one half-inch bump right of the far-left (0) position
            if ordered:
                active_ordered_num_id = add_list(
                    doc,
                    True,
                    items,
                    num_id=active_ordered_num_id,
                    left_twips=left_twips,
                )
            else:
                active_ordered_num_id = None
                add_list(doc, False, items, left_twips=left_twips)
        elif kind == "table":
            active_ordered_num_id = None
            add_table(doc, block[1])
            doc.add_paragraph("")
        elif kind == "image":
            # Keep active_ordered_num_id so numbering resumes after screenshots.
            add_image(doc, block[1])
        elif kind == "code":
            # Keep active_ordered_num_id so numbering resumes after code fences.
            add_code(doc, block[1])
        elif kind == "admonition":
            active_ordered_num_id = None
            add_admonition(doc, block[1], block[2])


ONLINE_NOTE = (
    "Note: the most recent version of this user guide may be online. "
    "Please visit docs.tenstorrent.com"
)

REVISION_PATHS = [
    *SOURCE_MD,
    GOLDEN,
    Path(__file__).resolve(),
]


def compute_revision() -> str:
    """Deterministic revision from git history of guide sources.

    Format: 1.<N> where N is the number of commits that touched the guide
    inputs. Same commit ⇒ same revision (safe for repeated CI rebuilds);
    content commits bump N automatically.
    """
    import subprocess

    try:
        result = subprocess.run(
            [
                "git",
                "-C",
                str(REPO),
                "rev-list",
                "--count",
                "HEAD",
                "--",
                *[str(p.relative_to(REPO)) if p.is_relative_to(REPO) else str(p) for p in REVISION_PATHS],
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode == 0 and result.stdout.strip().isdigit():
            return f"1.{int(result.stdout.strip())}"
    except Exception:  # noqa: BLE001
        pass
    return "1.0"


def prepare_front_matter(doc: Document) -> None:
    """Keep golden cover layout; refresh date/revision; ensure Blackhole® on the title."""
    today = date.today()
    revision = compute_revision()
    date_idx = None
    paras = list(doc.paragraphs)
    for i, para in enumerate(paras):
        text = para.text.strip()
        if re.match(r"^Revision\s+[\d.]+$", text, flags=re.IGNORECASE):
            set_runs_text(para, f"Revision {revision}")
        elif re.match(
            r"^(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},\s+\d{4}$",
            text,
        ):
            set_runs_text(para, today.strftime("%B %d, %Y"))
            date_idx = i
        elif text.startswith("Month, Date") or text == "Month, Date Year":
            set_runs_text(para, today.strftime("%B %d, %Y"))
            date_idx = i

    # Place the online-note directly under the revision date.
    if date_idx is not None:
        note_para = None
        for candidate in paras[date_idx + 1 : date_idx + 4]:
            text = candidate.text.strip()
            if text == "" or text.startswith("Note:"):
                note_para = candidate
                break
        if note_para is not None:
            set_runs_text(note_para, ONLINE_NOTE)
            for run in note_para.runs:
                if run.text:
                    run.italic = True
            # A little breathing room under the date line.
            note_para.paragraph_format.space_before = Pt(12)
            note_para.paragraph_format.space_after = Pt(0)

    # Title line 2 should read (Blackhole®) with a distinct ® run like QuietBox®.
    titles = [p for p in doc.paragraphs if p.style and p.style.name == "Title"]
    if len(titles) >= 2:
        para = titles[1]
        if "®" not in para.text:
            # Prefer inserting ® after the "Blackhole" run (mirrors QuietBox® title runs).
            for run in para.runs:
                if run.text.strip() == "Blackhole":
                    r_reg = para.add_run("®")
                    # Move ® run immediately after the Blackhole run in XML order.
                    run._element.addnext(r_reg._element)
                    r_reg.font.superscript = True
                    if run.font.size:
                        r_reg.font.size = run.font.size
                    if run.font.name:
                        r_reg.font.name = run.font.name
                    r_reg.bold = run.bold
                    break
            else:
                # Fallback: rewrite the whole subtitle.
                set_runs_text(para, "")
                append_text_with_trademarks(para, "(Blackhole®)")
def clear_chapter_body(doc: Document) -> None:
    """Keep cover + legal + trademarks + TOC heading; remove chapter content."""
    # Body tables live in the chapters — remove them all before cutting paragraphs.
    for table in list(doc.tables):
        delete_table(table)

    start_idx = None
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip() == "Specifications" and para.style and para.style.name.startswith("Heading"):
            start_idx = i
            break
    if start_idx is None:
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip() == "Table of Contents" and para.style and para.style.name.startswith("Heading"):
                # Keep TOC heading; clear everything after it
                start_idx = i + 1
                break
    if start_idx is None:
        return

    for para in list(doc.paragraphs[start_idx:]):
        delete_paragraph(para)


def build() -> Path:
    global _NUMBERING

    if GOLDEN.exists():
        template_path = GOLDEN
        TEMPLATE.write_bytes(GOLDEN.read_bytes())
    elif TEMPLATE.exists():
        template_path = TEMPLATE
    else:
        raise SystemExit(f"Missing golden/template DOCX under {SRC}")

    doc = Document(str(template_path))
    _NUMBERING = NumberingAllocator(doc)
    prepare_front_matter(doc)
    clear_chapter_body(doc)
    ensure_footer_page_numbers(doc)

    has_toc = any(
        p.text.strip() == "Table of Contents" and p.style and p.style.name.startswith("Heading")
        for p in doc.paragraphs
    )
    if not has_toc:
        doc.add_heading("Table of Contents", level=1)

    for idx, (filename, title) in enumerate(CHAPTERS):
        # Golden already has a page break after the TOC. Adding another before the
        # first chapter creates a blank page — only break before later chapters.
        if idx > 0:
            p = doc.add_paragraph()
            run = p.add_run()
            run.add_break(WD_BREAK.PAGE)
        doc.add_heading(title, level=1)

        md = clean_md((SRC / filename).read_text(encoding="utf-8"))
        md = re.sub(r"^# .+\n?", "", md, count=1, flags=re.MULTILINE)
        blocks = parse_blocks(md)
        render_blocks(doc, blocks, skip_first_h1=True)

    configure_toc_from_specifications(doc)
    superscript_all_registered(doc)

    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size // 1024} KB)")
    print(f"Style shell: {template_path.name}")
    return OUT_DOCX


def inputs_newer_than(target: Path) -> bool:
    """True if any build input is missing-target or newer than target."""
    if not target.exists():
        return True
    target_mtime = target.stat().st_mtime
    for path in BUILD_INPUTS:
        if path.exists() and path.stat().st_mtime > target_mtime:
            return True
    # Also rebuild PDF when the DOCX itself is newer.
    if target == OUT_PDF and OUT_DOCX.exists() and OUT_DOCX.stat().st_mtime > target_mtime:
        return True
    return False


def export_pdf_via_word(docx_path: Path, pdf_path: Path) -> None:
    """Export DOCX → PDF with Microsoft Word (best fidelity on macOS)."""
    import subprocess

    docx_path = docx_path.resolve()
    pdf_path = pdf_path.resolve()
    if pdf_path.exists():
        pdf_path.unlink()

    script = f'''
tell application "Microsoft Word"
  activate
  set docPath to POSIX file "{docx_path}"
  open docPath
  delay 2
  set theDoc to active document

  set n to count of fields of theDoc
  repeat with i from 1 to n
    try
      update field field i of theDoc
    end try
  end repeat
  try
    set sec to section 1 of theDoc
    set hf to get footer sec index header footer primary
    set tr to text object of hf
    repeat with j from 1 to (count of fields of tr)
      update field field j of tr
    end repeat
  end try

  set pdfPOSIX to "{pdf_path}"
  save as theDoc file name pdfPOSIX file format format PDF
  close theDoc saving no
end tell
'''
    result = subprocess.run(
        ["osascript", "-e", script],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0 or not pdf_path.exists():
        err = (result.stderr or result.stdout or "unknown error").strip()
        raise RuntimeError(f"Word PDF export failed: {err}")


def _find_libreoffice() -> str | None:
    import shutil

    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    for candidate in (
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/usr/local/bin/soffice",
        "/opt/homebrew/bin/soffice",
    ):
        if Path(candidate).exists():
            return candidate
    return None


def export_pdf_via_libreoffice(docx_path: Path, pdf_path: Path) -> None:
    """Export DOCX → PDF with LibreOffice (CI / headless fallback)."""
    import shutil
    import subprocess
    import tempfile

    soffice = _find_libreoffice()
    if not soffice:
        raise RuntimeError(
            "LibreOffice (soffice) not found. Install libreoffice-writer "
            "for headless PDF export."
        )

    docx_path = docx_path.resolve()
    pdf_path = pdf_path.resolve()
    pdf_path.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="qb2-pdf-") as tmp:
        tmp_dir = Path(tmp)
        cmd = [
            soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--norestore",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_dir),
            str(docx_path),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True, check=False)
        produced = tmp_dir / f"{docx_path.stem}.pdf"
        if result.returncode != 0 or not produced.exists():
            err = (result.stderr or result.stdout or "unknown error").strip()
            raise RuntimeError(f"LibreOffice PDF export failed: {err}")
        if pdf_path.exists():
            pdf_path.unlink()
        shutil.move(str(produced), str(pdf_path))


def export_pdf(docx_path: Path, pdf_path: Path) -> str:
    """Export PDF via Word when available, otherwise LibreOffice."""
    if (
        sys.platform == "darwin"
        and not os.environ.get("CI")
        and not os.environ.get("GITHUB_ACTIONS")
    ):
        try:
            export_pdf_via_word(docx_path, pdf_path)
            return "word"
        except Exception as word_exc:  # noqa: BLE001
            print(f"Word PDF export unavailable ({word_exc}); trying LibreOffice…")
    export_pdf_via_libreoffice(docx_path, pdf_path)
    return "libreoffice"


def build_pdf(docx_path: Path | None = None) -> Path:
    """Export the user guide PDF next to the DOCX."""
    docx_path = docx_path or OUT_DOCX
    if not docx_path.exists():
        raise SystemExit(f"Missing DOCX to export: {docx_path}")
    engine = export_pdf(docx_path, OUT_PDF)
    print(f"Wrote {OUT_PDF} ({OUT_PDF.stat().st_size // 1024} KB) via {engine}")
    return OUT_PDF


def in_ci() -> bool:
    return bool(os.environ.get("CI") or os.environ.get("GITHUB_ACTIONS"))


def ensure_user_guide(force: bool = False) -> Path | None:
    """Rebuild DOCX/PDF when sources change. Returns PDF path if available.

    In CI, always rebuild so deploys never depend on a human regenerating the PDF.
    """
    force = force or in_ci()
    need_docx = force or inputs_newer_than(OUT_DOCX)
    need_pdf = force or inputs_newer_than(OUT_PDF)

    if need_docx:
        build()
        need_pdf = True
    elif need_pdf and not OUT_DOCX.exists():
        build()

    if need_pdf or not OUT_PDF.exists():
        try:
            return build_pdf()
        except Exception as exc:  # noqa: BLE001
            if in_ci():
                raise RuntimeError(
                    f"QuietBox 2 user-guide PDF export failed in CI: {exc}"
                ) from exc
            if OUT_PDF.exists():
                print(f"PDF export failed ({exc}); keeping existing {OUT_PDF.name}")
                return OUT_PDF
            print(f"PDF export failed: {exc}")
            return None
    return OUT_PDF if OUT_PDF.exists() else None


if __name__ == "__main__":
    force = "--force" in sys.argv
    pdf_only = "--pdf-only" in sys.argv
    if pdf_only:
        build_pdf()
    else:
        ensure_user_guide(force=force)
